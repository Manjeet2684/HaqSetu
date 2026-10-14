"""
Generate HaqSetu submission deliverables:
  - deliverables/HaqSetu_Lean_Canvas.pdf  (reportlab)
  - deliverables/HaqSetu_Concept_Note.docx (python-docx)

NOTE: A PowerPoint file is NOT generated because no PPTX dependency is installed.
An editable slide-by-slide presentation outline is provided in
deliverables/HaqSetu_Presentation_Outline.md as the authoritative slide source.
Screenshot capture from the running application remains a manual task.
"""

from __future__ import annotations

from pathlib import Path


def generate_lean_canvas_pdf(output_path: Path) -> None:
    """Generate the 9-block HaqSetu Lean Canvas PDF using pre-installed reportlab."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CanvasTitle",
        parent=styles["Heading1"],
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#0F4C5C"),
        spaceAfter=6,
    )
    sub_style = ParagraphStyle(
        "CanvasSub",
        parent=styles["Normal"],
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#566270"),
        spaceAfter=10,
    )
    cell_title = ParagraphStyle(
        "CellTitle",
        parent=styles["Heading3"],
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor("#0F4C5C"),
        spaceAfter=4,
    )
    cell_body = ParagraphStyle(
        "CellBody",
        parent=styles["Normal"],
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#1B2733"),
    )

    def block(heading: str, bullets: list[str]) -> list[Paragraph]:
        text = "<br/>".join(f"&bull; {b}" for b in bullets)
        return [Paragraph(f"<b>{heading}</b>", cell_title), Paragraph(text, cell_body)]

    row1 = [
        block(
            "1. Problem",
            [
                "Eligible students miss scholarships due to fragmented PDFs and unclear criteria.",
                "Generic chatbots hallucinate income limits, marks cutoffs, and deadlines.",
                "Students lack a clear prerequisite order for certificates and OTR registration.",
            ],
        ),
        block(
            "4. Solution (Implemented MVP)",
            [
                "Deterministic 3-valued Kleene eligibility engine (ELIGIBLE / INELIGIBLE / NEEDS_VERIFICATION).",
                "DAG Action Planner (graphlib.TopologicalSorter) with DONE/READY/BLOCKED/GATED states.",
                "Stdlib TF-IDF policy clause retriever with strict refusal when evidence is unverified.",
            ],
        ),
        block(
            "3. Unique Value Proposition",
            [
                "Eligibility Condition -> Supporting Evidence -> Decision -> Next Action.",
                "Zero hallucinated verdicts: unverified schemes fail closed to LISTING_ONLY.",
                "Every evaluated condition links to its exact clause, section, and source URL.",
            ],
        ),
        block(
            "9. Unfair Advantage",
            [
                "Strict structural separation between SchemeReadiness (ASSESSABLE vs LISTING_ONLY) "
                "and ApplicantOutcome.",
                "Safety gates cannot be bypassed by marking earlier steps complete.",
            ],
        ),
        block(
            "2. Customer Segments",
            [
                "Primary: Indian school and college students seeking public scholarships (SDG 4/1/10).",
                "Secondary (Hypothesis): School/college scholarship cells, CSCs, NGO facilitators.",
            ],
        ),
    ]

    row2 = [
        block(
            "8. Key Metrics (Actual Implemented)",
            [
                "Automated pytest suite: 32 unit & integration tests across engine, planner, retriever, API.",
                "Retrieval benchmark (10 synthetic queries): Top-1=1.0, Top-3=1.0, MRR=1.0, "
                "Refusal=1.0 (small synthetic benchmark — not independent validation).",
                "Zero external ML/LLM dependencies; runs locally on Python 3.14.",
            ],
        ),
        block(
            "5. Channels (Hypothesis)",
            [
                "Local offline-capable institutional kiosk deployment.",
                "Open-source educational utility for college nodal cells and NGO field workers.",
            ],
        ),
        block(
            "7. Cost Structure (Actual MVP)",
            [
                "Zero cloud hosting, database, or paid LLM API costs.",
                "Runs locally on standard student laptop (Python 3.14, FastAPI, Pydantic v2, vanilla JS).",
            ],
        ),
        block(
            "6. Sustainability / Revenue (Hypothesis)",
            [
                "Public-good open-source tool aligned with SDG 4.",
                "Potential CSR/NGO grant sponsorship for human policy verification cycles (Hypothesis).",
            ],
        ),
        block(
            "Honest Scope & Verification Note",
            [
                "Real candidates (PM-USP, UP Pre-Matric) remain LISTING_ONLY until human verification "
                "of current cycle guidelines.",
                "Interactive assessment demonstrated safely via explicitly labelled FICTIONAL_FIXTURE.",
                "No RAM footprint claim published — not measured in an automated test.",
            ],
        ),
    ]

    table = Table([row1, row2], colWidths=[155, 155, 155, 155, 155])
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.75, colors.HexColor("#E4DED3")),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FAF7F2")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )

    story = [
        Paragraph("HaqSetu — AI Welfare and Scholarship Navigator (Lean Canvas)", title_style),
        Paragraph(
            "Primary SDG: SDG 4 (Quality Education) | Secondary: SDG 1 (No Poverty), "
            "SDG 10 (Reduced Inequalities) | IBM SkillsBuild Internship Deliverable",
            sub_style,
        ),
        Spacer(1, 4),
        table,
    ]
    doc.build(story)


def generate_concept_note_docx(output_path: Path) -> None:
    """Generate the HaqSetu Concept Note DOCX using pre-installed python-docx."""
    from docx import Document

    doc = Document()
    doc.add_heading("HaqSetu — AI Welfare and Scholarship Navigator: Concept Note", level=0)
    doc.add_paragraph(
        "Primary SDG: SDG 4 (Quality Education) | Related SDGs: SDG 1 (No Poverty), "
        "SDG 10 (Reduced Inequalities)"
    )

    sections = [
        (
            "1. Problem Statement & Vision",
            "Governments publish welfare and scholarship schemes, yet intended beneficiaries struggle to "
            "discover relevant opportunities, interpret complex eligibility clauses, prepare prerequisite "
            "documents in the right order, or locate official grievance channels. HaqSetu bridges this gap "
            "through the workflow: Eligibility Condition -> Supporting Evidence -> Decision -> Next Action.",
        ),
        (
            "2. Implemented Architecture & Technical Honesty",
            "HaqSetu is built as a lightweight local modular monolith using Python 3.14.5, FastAPI, Pydantic v2, "
            "Python's standard-library graphlib.TopologicalSorter, and a responsive Civic Ledger "
            "HTML/CSS/vanilla JS interface. It separates SchemeReadiness (ASSESSABLE vs LISTING_ONLY) from "
            "ApplicantOutcome (ELIGIBLE, INELIGIBLE, NEEDS_VERIFICATION). The full test suite runs 32 pytest "
            "cases (P2/P3: 18, P5: 5, P4/P7: 9 total across 4 modules, including parametrized cases).",
        ),
        (
            "3. Honest AI/ML Role & Evaluation",
            "Eligibility decisions are never delegated to a machine-learning classifier or generative LLM. "
            "Instead, HaqSetu uses symbolic three-valued Kleene logic for deterministic eligibility evaluation, "
            "DAG topological sorting for action planning, and a pure-Python standard-library TF-IDF + cosine "
            "similarity retriever for source-grounded policy Q&A (/api/ask). A 10-query synthetic smoke "
            "benchmark achieved Top-1=1.0, Top-3=1.0, MRR=1.0, Refusal=1.0. These are small synthetic "
            "figures — not independently validated accuracy claims.",
        ),
        (
            "4. Policy Verification Safeguards & Current Status",
            "Real scheme candidates in data/schemes.json (PM-USP Central Sector Scholarship on "
            "scholarships.gov.in and Uttar Pradesh Pre-Matric Scholarship on scholarship.up.gov.in) are "
            "maintained as LISTING_ONLY until a human reviewer inspects current academic-year guidelines and "
            "records verbatim clauses. Interactive eligibility and planner mechanics are demonstrated safely "
            "via an isolated fictional fixture (demo-assessable) that carries persistent FICTIONAL DEMO "
            "WARNING banners. No eligibility rules, income limits, marks thresholds, or deadlines have been "
            "fabricated for real schemes.",
        ),
        (
            "5. Privacy, Security & Limitations",
            "HaqSetu collects no Aadhaar numbers, bank details, or logins. User answers remain stateless "
            "and are never logged. The application enforces a 32 KB payload cap and defensive security "
            "headers. Future scope (not claimed as implemented in the MVP) includes multi-state pin-code "
            "discovery, vernacular translation, and nodal-officer SLA tracking. No RAM footprint claim is "
            "published because it was not measured in an automated test.",
        ),
    ]

    for heading, body in sections:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(body)

    doc.save(str(output_path))


def generate_presentation_outline_md(output_path: Path) -> None:
    """
    Generate an editable slide-by-slide Markdown source for the presentation.
    NOTE: A PPTX package is not installed; this Markdown outline is the primary
    presentation deliverable. Screenshots from the running application should be
    captured manually and inserted where noted.
    """
    content = """\
# HaqSetu — Presentation Outline (Slide-by-Slide Source)

> **Note:** This Markdown file is the editable slide source. No PowerPoint generator
> is installed in this environment. Use this outline to build slides manually in
> PowerPoint, Google Slides, or Keynote. Screenshot placeholders must be filled
> by capturing the running application at `http://127.0.0.1:8000`.

---

## Slide 1 — Title Slide
- **Product Name:** HaqSetu — AI Welfare and Scholarship Navigator
- **Tagline:** "The opportunity exists. The person who needs it exists. The bridge connecting them is often missing."
- **SDGs:** Primary: SDG 4 (Quality Education) | Secondary: SDG 1, SDG 10
- **IBM SkillsBuild Internship Deliverable**

---

## Slide 2 — Problem Statement
- Students miss scholarships due to fragmented PDFs and unclear eligibility criteria
- Chatbots hallucinate income limits, marks cutoffs, and deadlines
- No clear prerequisite order for certificates, OTR registration, nodal officer verification
- Real, unverified schemes must not be turned into misleading automated verdicts

---

## Slide 3 — HaqSetu Workflow
```
Eligibility Condition → Supporting Evidence → Decision → Next Action
```
- Discovery: Browse real scheme catalogue with verification status
- Assessment: Three-valued Kleene logic (ELIGIBLE / INELIGIBLE / NEEDS_VERIFICATION)
- Action Planning: Dependency-aware DAG with DONE/READY/BLOCKED/GATED
- Q&A: Source-grounded TF-IDF retrieval with explicit refusal

---

## Slide 4 — Architecture (Modular Monolith)
- `app/schemas.py` → Pydantic v2 contracts + fail-closed loaders
- `app/engine.py` → Kleene 3-valued eligibility evaluation
- `app/planner.py` → graphlib.TopologicalSorter action planner
- `app/retriever.py` → Stdlib TF-IDF + cosine similarity
- `app/catalog.py` → Real/demo separation catalogue service
- `app/main.py` → FastAPI routes + security middleware + static mount
- `js/haqsetu_api_client.js` → Progressive P6 API bridge (safe DOM)

---

## Slide 5 — Safety by Design
- **LISTING_ONLY** real schemes → `applicant_outcome = null` always
- **Safety gates** cannot be bypassed by marking actions complete
- **Failed global mandatory condition** cannot be overridden by passing pathway
- **Unknown/missing answers** never treated as passing
- **Fictional demo** always labelled; never leaks to real endpoints
- **No Aadhaar, bank data, or PII** collected or logged

---

## Slide 6 — Scheme Catalogue (Screenshot)
> [PLACEHOLDER: Screenshot of real scheme catalogue panel from http://127.0.0.1:8000]
- Shows PM-USP and UP Pre-Matric as LISTING_ONLY
- Human verification checklist visible
- Official portal, helpdesk, and CPGRAMS links

---

## Slide 7 — Demo Assessment (Screenshot)
> [PLACEHOLDER: Screenshot of fictional demo assessment from http://127.0.0.1:8000]
- FICTIONAL DEMO WARNING banner persistent
- Condition-by-condition outcomes (SATISFIED/FAILED/UNRESOLVED)
- GATED submission when eligibility unresolved
- ELIGIBLE outcome when all conditions met

---

## Slide 8 — Policy Q&A Retrieval (Screenshot)
> [PLACEHOLDER: Screenshot of policy Q&A panel from http://127.0.0.1:8000]
- Real scheme query → INSUFFICIENT_VERIFIED_EVIDENCE refusal
- Demo scheme query → source-linked fictional clause
- Source title, section, cycle, score, and official URL visible

---

## Slide 9 — Test Suite & Benchmark
- **32 pytest cases** collected and passing (Python 3.14.5, pytest 9.1.1)
  - 11 engine/schema, 7 planner (incl. 3 parametrized), 5 retrieval, 9 API/security
- **10-query synthetic smoke benchmark** (data/eval_queries.json):
  - Top-1 Accuracy: 1.0 | Top-3 Recall: 1.0 | MRR: 1.0 | Refusal Accuracy: 1.0
  - ⚠️ Small synthetic corpus — not independent validation

---

## Slide 10 — Honest Current Status & Roadmap
**Implemented:**
- Full P1–P8 pipeline: engine, planner, retrieval, API, UI bridge
- 2 real scheme stubs (LISTING_ONLY, safe by design)
- 1 verified assessable fictional demo scenario

**Pending human action:**
- Inspect official government portals for current academic-year rules
- Record verbatim clauses, section/page, reviewer name, date
- Only then promote a scheme from LISTING_ONLY → ASSESSABLE

**Future scope (not in MVP):**
- Vernacular translation, pin-code district lookup, multi-scheme comparison

---
"""
    output_path.write_text(content, encoding="utf-8")


if __name__ == "__main__":
    out_dir = Path(__file__).resolve().parent.parent / "deliverables"
    out_dir.mkdir(parents=True, exist_ok=True)
    pdf_file = out_dir / "HaqSetu_Lean_Canvas.pdf"
    docx_file = out_dir / "HaqSetu_Concept_Note.docx"
    md_file = out_dir / "HaqSetu_Presentation_Outline.md"

    generate_lean_canvas_pdf(pdf_file)
    print(f"Generated: {pdf_file}")

    generate_concept_note_docx(docx_file)
    print(f"Generated: {docx_file}")

    generate_presentation_outline_md(md_file)
    print(f"Generated: {md_file}")
